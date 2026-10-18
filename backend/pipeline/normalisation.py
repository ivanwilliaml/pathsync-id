"""Evaluate embedding normalisation against human labels (plan A3, paper III.D level 1).

    python -m pipeline.normalisation sample            # 300 phrases, 100 per similarity band (locked)
    python -m pipeline.normalisation sheet <name>      # Excel for one labeller (no system prediction or band shown)
    python -m pipeline.normalisation import <file>...  # labelled .xlsx or .csv -> data/seed/normalisation_gold.json
    python -m pipeline.normalisation eval              # curve, purity, chosen threshold -> normalisation_eval.json

Each labeller writes, per phrase, the correct taxonomy skill id or "none" (not a skill, or a skill
the taxonomy does not have). Phrases where the two labellers disagree are left out of the gold until
a third rater fills `adjudicated` in the gold file.

Threshold rule (written before the labels exist): the auto band starts at the lowest threshold whose
precision on the gold is >= 0.95. As in the live gates, a phrase above the threshold with no word in
common with the matched skill form (lexical_support) is not mapped automatically.
"""
from __future__ import annotations

import csv
import json
import random
import sys
from collections import Counter, defaultdict
from pathlib import Path

from pathsync.config import MODELS_DIR, PROCESSED_DIR, SEED_DIR, SIM_AUTO, SIM_REVIEW
from pathsync.gold import wilson

SAMPLE = SEED_DIR / "normalisation_sample.json"
GOLD = SEED_DIR / "normalisation_gold.json"
REPORT = MODELS_DIR / "normalisation_eval.json"
BANDS = (("below_0.70", 0.0, SIM_REVIEW), ("0.70_to_0.85", SIM_REVIEW, SIM_AUTO), ("from_0.85", SIM_AUTO, 1.01))
PER_BAND = 100
AUTO_PRECISION_TARGET = 0.95
THRESHOLDS = [round(0.60 + 0.05 * i, 2) for i in range(8)]       # 0.60 .. 0.95
SHEET_GUIDE = [
    "Untuk setiap frasa, pilih di kolom kuning id skill yang dimaksud frasa itu (dropdown), atau 'none'.",
    "'none' berarti: bukan skill (benefit, sifat, nama perusahaan, domain bisnis tanpa kemampuan), soft skill, atau skill yang tidak ada di daftar.",
    "Baca juga kalimat asalnya. Contoh: 'menjalankan campaign iklan berbayar' -> social_media_ads atau search_ads, sesuai platform yang disebut.",
    "Kalau dua skill sama-sama cocok, pilih yang paling spesifik dan tulis yang lain di kolom catatan.",
    "Kerjakan sendiri, jangan berdiskusi dengan pelabel lain sampai keduanya selesai. Sekitar 1 jam untuk 300 frasa.",
    "Setelah selesai, kirim berkas ini kembali tanpa mengubah kolom id.",
]


def cmd_sample(seed: int = 20261015) -> None:
    from sqlalchemy import select

    from pathsync.db import PostingMention, SessionLocal

    if SAMPLE.exists():
        raise SystemExit(f"{SAMPLE.name} exists; the sample is locked")
    with SessionLocal() as s:
        rows = s.execute(select(PostingMention.id, PostingMention.surface, PostingMention.quote, PostingMention.similarity,
                                PostingMention.status).where(PostingMention.method == "pattern")).all()
    seen, uniq = set(), []
    for r in rows:                                  # one row per distinct phrase (case-insensitive)
        k = r.surface.strip().lower()
        if k and k not in seen:
            seen.add(k)
            uniq.append(r)
    rng = random.Random(seed)
    picked = []
    for name, lo, hi in BANDS:
        band = [r for r in uniq if lo <= r.similarity < hi]
        for r in rng.sample(band, min(PER_BAND, len(band))):
            picked.append({"id": f"n{len(picked) + 1:03d}", "surface": r.surface, "quote": r.quote[:300], "band": name})
    SAMPLE.write_text(json.dumps({"seed": seed, "note": "Locked sample for plan A3; phrases only, no system output.",
                                  "phrases": picked}, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"{len(picked)} phrases -> {SAMPLE}")


LABEL_COL = "label (skill id or none)"


def cmd_sheet(name: str) -> None:
    """Excel sheet for one labeller: phrases in a shuffled order (the similarity band is never shown), a dropdown
    of every taxonomy skill id plus "none", and the skill list with names and example words."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    from pathsync.taxonomy import load_skills

    doc = json.loads(SAMPLE.read_text(encoding="utf-8"))
    phrases = sorted(doc["phrases"], key=lambda p: random.Random(f"{doc['seed']}|{p['id']}").random())
    skills = sorted(load_skills().items(), key=lambda kv: (kv[1].cluster, kv[0]))
    wb = Workbook()
    guide = wb.active
    guide.title = "Petunjuk"
    for line in [f"Label normalisasi skill, pelabel: {name}", *SHEET_GUIDE]:
        guide.append([line])
    guide["A1"].font = Font(bold=True, size=13)
    guide.column_dimensions["A"].width = 140
    sk = wb.create_sheet("Daftar skill")
    sk.append(["id", "nama", "klaster", "contoh kata di lowongan"])
    for sid, skill in skills:
        sk.append([sid, skill.name_id or skill.name, skill.cluster, ", ".join(skill.aliases[:6])])
    sk.append(["none", "bukan skill atau tidak ada di daftar", "", ""])
    for col, w in zip("ABCD", (30, 40, 18, 70), strict=True):
        sk.column_dimensions[col].width = w
    sk.freeze_panes = "B2"
    ws = wb.create_sheet("Frasa", 1)
    ws.append(["no", "id", "frasa", "kalimat asal di lowongan", LABEL_COL, "catatan"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A5F")
    for i, p in enumerate(phrases, 1):
        ws.append([i, p["id"], p["surface"], p["quote"], "", ""])
    dv = DataValidation(type="list", formula1=f"='Daftar skill'!$A$2:$A${len(skills) + 2}", allow_blank=True,
                        showErrorMessage=True, errorTitle="Id tidak dikenal", error="Pilih id dari daftar atau tulis none.")
    ws.add_data_validation(dv)
    dv.add(f"E2:E{len(phrases) + 1}")
    for col, w in zip("ABCDEF", (5, 8, 36, 90, 30, 30), strict=True):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        row[4].fill = PatternFill("solid", fgColor="FFF5CC")
    ws.freeze_panes = "E2"
    out = PROCESSED_DIR / f"normalisation_sheet_{name}.xlsx"
    wb.save(out)
    print(f"sheet for {name}: {out}")


def _read_labels(path: str) -> list[tuple[str, str]]:
    if path.lower().endswith(".xlsx"):
        from openpyxl import load_workbook

        ws = load_workbook(path, read_only=True)["Frasa"]
        return [(str(r[1]), str(r[4] or "").strip()) for r in ws.iter_rows(min_row=2, values_only=True) if r[1]]
    with open(path, encoding="utf-8-sig") as f:
        return [(row["id"], (row.get(LABEL_COL) or "").strip()) for row in csv.DictReader(f)]


def cmd_import(paths: list[str]) -> None:
    from pathsync.taxonomy import load_skills

    valid = set(load_skills()) | {"none"}
    labels: dict[str, dict[str, str]] = defaultdict(dict)
    for p in paths:
        who = Path(p).stem.replace("normalisation_sheet_", "")
        for pid, lab in _read_labels(p):
            if lab:
                if lab not in valid:
                    raise SystemExit(f"{p}: unknown label '{lab}' on {pid}")
                labels[pid][who] = lab
    old = json.loads(GOLD.read_text(encoding="utf-8")) if GOLD.exists() else {"phrases": {}}
    phrases = {}
    for pid, by in labels.items():
        prev = old["phrases"].get(pid, {})
        phrases[pid] = {"labels": by, "adjudicated": prev.get("adjudicated")}
    GOLD.write_text(json.dumps({"note": "Two human labellers per phrase (plan A3); 'adjudicated' resolves disagreements.",
                                "phrases": phrases}, ensure_ascii=False, indent=1), encoding="utf-8")
    both = [v for v in phrases.values() if len(v["labels"]) >= 2]
    agree = sum(len(set(v["labels"].values())) == 1 for v in both)
    print(f"{len(phrases)} phrases labelled, {len(both)} by two people, raw agreement {agree}/{len(both)}, "
          f"kappa {labeller_kappa()['kappa']}")


def gold_labels() -> dict[str, str]:
    """phrase id -> agreed label (or the adjudicated one)."""
    doc = json.loads(GOLD.read_text(encoding="utf-8"))
    out = {}
    for pid, v in doc["phrases"].items():
        labs = set(v["labels"].values())
        if v.get("adjudicated"):
            out[pid] = v["adjudicated"]
        elif len(v["labels"]) >= 2 and len(labs) == 1:
            out[pid] = labs.pop()
    return out


def cohen_kappa(a: list[str], b: list[str]) -> float | None:
    n = len(a)
    if not n:
        return None
    po = sum(x == y for x, y in zip(a, b, strict=True)) / n
    ca, cb = Counter(a), Counter(b)
    pe = sum(ca[k] * cb.get(k, 0) for k in ca) / (n * n)
    return round((po - pe) / (1 - pe), 4) if pe < 1 else 1.0


def labeller_kappa() -> dict:
    """Cohen's kappa over the label each of the first two labellers gave to the same phrase."""
    doc = json.loads(GOLD.read_text(encoding="utf-8"))
    names = sorted({n for v in doc["phrases"].values() for n in v["labels"]})[:2]
    both = [v["labels"] for v in doc["phrases"].values() if all(n in v["labels"] for n in names)]
    return {"labellers": names, "n": len(both),
            "kappa": cohen_kappa([b[names[0]] for b in both], [b[names[1]] for b in both]) if len(names) == 2 else None}


def evaluate(gold: dict[str, str], predictions: dict[str, tuple]) -> dict:
    """predictions: phrase id -> (nearest skill, similarity[, lexically supported]) from the current index.
    An unsupported phrase is never auto-mapped, whatever its similarity (same rule as the live gates)."""
    ids = [i for i in gold if i in predictions]
    skill_phrases = [i for i in ids if gold[i] != "none"]
    supported = {i: (predictions[i][2] if len(predictions[i]) > 2 else True) for i in ids}
    curve = []
    for t in THRESHOLDS:
        auto = [i for i in ids if predictions[i][1] >= t and supported[i]]
        right = sum(predictions[i][0] == gold[i] for i in auto)
        below = [i for i in ids if not (predictions[i][1] >= t and supported[i])]
        curve.append({"threshold": t, "auto": len(auto), "precision": round(right / len(auto), 4) if auto else None,
                      "precision_ci95": wilson(right, len(auto)),
                      "coverage": round(right / len(skill_phrases), 4) if skill_phrases else None,
                      "non_skill_kept_out": round(sum(gold[i] == "none" for i in below) /
                                                  max(1, sum(gold[i] == "none" for i in ids)), 4)})
    ok = [c for c in curve if c["precision"] is not None and c["precision"] >= AUTO_PRECISION_TARGET]
    chosen = min((c["threshold"] for c in ok), default=None)
    groups: dict[str, Counter] = defaultdict(Counter)                      # cluster purity (paper III.D)
    for i in ids:
        groups[predictions[i][0]][gold[i]] += 1
    n = sum(sum(c.values()) for c in groups.values())
    purity = sum(c.most_common(1)[0][1] for c in groups.values()) / n if n else None
    band_now = [c for c in curve if c["threshold"] == round(SIM_AUTO, 2)]
    return {"n_gold": len(ids), "n_skill_phrases": len(skill_phrases), "curve": curve,
            "rule": f"lowest threshold with precision >= {AUTO_PRECISION_TARGET}", "chosen_threshold": chosen,
            "current_threshold": SIM_AUTO, "current_auto_precision": band_now[0]["precision"] if band_now else None,
            "cluster_purity": round(purity, 4) if purity is not None else None}


def cmd_eval() -> None:
    from pathsync.extraction import skill_index

    if not GOLD.exists():
        raise SystemExit("no normalisation_gold.json yet: run sample, sheet, label, import")
    phrases = {p["id"]: p for p in json.loads(SAMPLE.read_text(encoding="utf-8"))["phrases"]}
    gold = gold_labels()
    ids = [i for i in gold if i in phrases]
    from pathsync.extraction.base import MIN_LEXICAL_SUPPORT, lexical_support

    near = skill_index().nearest([phrases[i]["surface"] for i in ids])
    preds = {i: (sid, float(sim), lexical_support(phrases[i]["surface"], form) >= MIN_LEXICAL_SUPPORT)
             for i, (sid, sim, form) in zip(ids, near, strict=True)}
    out = {**evaluate(gold, preds), "labellers": labeller_kappa()}
    REPORT.write_text(json.dumps(out, indent=1), encoding="utf-8")
    print(json.dumps({k: out[k] for k in ("n_gold", "chosen_threshold", "current_auto_precision", "cluster_purity",
                                          "labellers")}))


def main() -> None:
    cmd, *rest = sys.argv[1:] or ["eval"]
    if cmd == "sample":
        cmd_sample()
    elif cmd == "sheet" and rest:
        cmd_sheet(rest[0])
    elif cmd == "import" and rest:
        cmd_import(rest)
    elif cmd == "eval":
        cmd_eval()
    else:
        raise SystemExit(__doc__)


if __name__ == "__main__":
    main()
