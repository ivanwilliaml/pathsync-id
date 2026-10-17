"""Human gold set workflow (plan A1).

    python -m pipeline.gold split       # writes data/seed/gold_split.json once (refuses to overwrite)
    python -m pipeline.gold split-bisnis   # writes data/seed/gold_split_bisnis.json once (6 business roles, 5 each)
    python -m pipeline.gold status      # progress per annotator on the split
    python -m pipeline.gold consensus   # annotations -> gold_consensus.json + open disputes
    python -m pipeline.gold sheet <name> [out.xlsx]        # Excel sheet for one annotator (offline)
    python -m pipeline.gold import-sheet <name> <xlsx>     # filled sheet -> annotations in the database

Annotators label in Admin > Annotation, or offline in an Excel sheet (same postings, same order, posting
text only, never the extractor output); postings in the split are served first. Disputed
(posting, skill) pairs go to a third rater, whose decisions are written to
data/seed/gold_adjudication.json as {"posting_id": {"skill_id": true|false}}.
"""
from __future__ import annotations

import json
import sys
from collections import defaultdict

from sqlalchemy import select

from pathsync.config import MODELS_DIR
from pathsync.db import Annotation, Posting, SessionLocal, init_db
from pathsync.gold import (
    ADJUDICATION_FILE,
    CONSENSUS_FILE,
    EXTRA_SPLIT_FILE,
    EXTRA_SPLIT_SEED,
    SPLIT_FILE,
    agreement,
    consensus,
    load_split,
    load_split_all,
    make_split,
)
from pathsync.taxonomy import load_skills


def _annotations(s, ids: set[str]) -> dict[str, dict[str, set]]:
    by_post: dict[str, dict[str, set]] = defaultdict(dict)
    for a in s.scalars(select(Annotation).where(Annotation.posting_id.in_(ids))).all():
        by_post[a.posting_id][a.annotator] = set(a.skill_ids)
    return by_post


def cmd_split() -> None:
    if SPLIT_FILE.exists():
        raise SystemExit(f"{SPLIT_FILE.name} already exists; the split is locked. Delete it by hand only "
                         "if no annotation has started.")
    with SessionLocal() as s:
        rows = s.execute(select(Posting.id, Posting.role).where(Posting.duplicate_of.is_(None))).all()
    doc = make_split(dict(rows))
    SPLIT_FILE.write_text(json.dumps(doc, indent=1), encoding="utf-8")
    print(f"dev {len(doc['dev'])}, test {len(doc['test'])} -> {SPLIT_FILE}")


def cmd_split_bisnis() -> None:
    from pipeline.add_business_roles import BUSINESS_ROLES

    if EXTRA_SPLIT_FILE.exists():
        raise SystemExit(f"{EXTRA_SPLIT_FILE.name} already exists; the split is locked.")
    main = load_split() or {"dev": [], "test": []}
    taken = set(main["dev"]) | set(main["test"])
    with SessionLocal() as s:
        rows = s.execute(select(Posting.id, Posting.role).where(Posting.duplicate_of.is_(None),
                                                                Posting.role.in_(BUSINESS_ROLES))).all()
    doc = make_split({pid: role for pid, role in rows if pid not in taken}, seed=EXTRA_SPLIT_SEED)
    doc["note"] = "Locked supplement to gold_split.json for the six business roles (plan A9). Same rules: tune on dev only."
    EXTRA_SPLIT_FILE.write_text(json.dumps(doc, indent=1), encoding="utf-8")
    print(f"dev {len(doc['dev'])}, test {len(doc['test'])} -> {EXTRA_SPLIT_FILE}")


def cmd_status() -> None:
    split = load_split_all() or {"dev": [], "test": []}
    ids = set(split["dev"]) | set(split["test"])
    with SessionLocal() as s:
        by_post = _annotations(s, ids)
    per = defaultdict(int)
    for anns in by_post.values():
        for n in anns:
            per[n] += 1
    print(json.dumps({"split_postings": len(ids), "annotated_by": dict(per),
                      "with_two_or_more": sum(len(a) >= 2 for a in by_post.values())}, indent=1))


def cmd_consensus() -> None:
    split = load_split_all()
    if not split:
        raise SystemExit("run `python -m pipeline.gold split` first")
    ids = set(split["dev"]) | set(split["test"])
    adjudication = json.loads(ADJUDICATION_FILE.read_text(encoding="utf-8")) if ADJUDICATION_FILE.exists() else {}
    with SessionLocal() as s:
        by_post = _annotations(s, ids)
        roles = dict(s.execute(select(Posting.id, Posting.role).where(Posting.id.in_(ids))).all())
    doc, open_disputes = consensus(by_post, roles, adjudication)
    agree = agreement(by_post, list(load_skills()))
    CONSENSUS_FILE.write_text(json.dumps(doc, indent=1), encoding="utf-8")
    (MODELS_DIR / "gold_agreement.json").write_text(json.dumps({"pairs": agree, "open_disputes": open_disputes},
                                                               indent=1), encoding="utf-8")
    print(f"consensus postings: {len(doc['postings'])}, open disputes: {len(open_disputes)}")
    for p in agree:
        print(f"kappa {p['annotators']}: {p['kappa']} on {p['postings']} postings (target >= 0.70)")


GUIDE = [
    "Isi kolom 'skill' dengan id skill dari sheet 'Daftar skill', dipisah koma. Contoh: sql, excel, requirements_elicitation",
    "Hitung: syarat atau tugas yang tertulis eksplisit, daftar alternatif ('Python, R, atau sejenis'), dan nilai tambah ('diutamakan').",
    "Jangan hitung: bidang gelar, nama perusahaan atau produknya, domain bisnis tanpa kemampuan, soft skill, benefit.",
    "Baca seluruh teks lowongan. Kerjakan sendiri, jangan lihat lembar anotator lain atau halaman Job Market.",
    "Lowongan tanpa skill dari daftar: tulis 'none'. Ragu antara dua skill: pilih yang paling spesifik dan tulis di kolom catatan.",
    "Panduan lengkap: docs/validasi/PANDUAN_ANOTASI.md. Setelah selesai, kirim berkas ini kembali tanpa mengubah kolom id.",
]


def cmd_sheet(name: str, out: str | None = None) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    from pathsync.config import PROCESSED_DIR

    split = load_split_all()
    if not split:
        raise SystemExit("run `python -m pipeline.gold split` first")
    order = split["dev"] + split["test"]
    with SessionLocal() as s:
        posts = {p.id: p for p in s.scalars(select(Posting).where(Posting.id.in_(order))).all()}
    wb = Workbook()
    guide = wb.active
    guide.title = "Petunjuk"
    guide.append([f"Anotasi skill lowongan, anotator: {name}"])
    guide["A1"].font = Font(bold=True, size=13)
    for line in GUIDE:
        guide.append([line])
    guide.column_dimensions["A"].width = 140
    ws = wb.create_sheet("Lowongan")
    ws.append(["no", "id", "peran", "judul", "teks lowongan", "skill", "catatan"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A5F")
    for i, pid in enumerate(order, 1):
        p = posts[pid]
        ws.append([i, pid, p.role, p.title, p.description[:32000], "", ""])
    for col, w in zip("ABCDEFG", (5, 26, 18, 30, 110, 40, 30), strict=True):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        row[5].fill = PatternFill("solid", fgColor="FFF5CC")
    ws.freeze_panes = "F2"
    sk = wb.create_sheet("Daftar skill")
    sk.append(["id", "nama", "nama (EN)", "klaster", "contoh kata di lowongan"])
    for c in sk[1]:
        c.font = Font(bold=True)
    for sid, skill in sorted(load_skills().items(), key=lambda kv: (kv[1].cluster, kv[0])):
        sk.append([sid, skill.name_id or skill.name, skill.name_en, skill.cluster, ", ".join(skill.aliases[:6])])
    for col, w in zip("ABCDE", (30, 34, 34, 18, 70), strict=True):
        sk.column_dimensions[col].width = w
    sk.freeze_panes = "B2"
    path = out or str(PROCESSED_DIR / f"anotasi_lowongan_{name}.xlsx")
    wb.save(path)
    print(f"{len(order)} postings -> {path}")


def cmd_import_sheet(name: str, path: str) -> None:
    from openpyxl import load_workbook

    valid = set(load_skills())
    ws = load_workbook(path, read_only=True)["Lowongan"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    got, empty, bad = {}, 0, []
    for _, pid, _role, _title, _text, skills, *_ in rows:
        raw = str(skills or "").strip()
        if not raw:
            empty += 1
            continue
        ids = [] if raw.lower() == "none" else [x.strip() for x in raw.replace(";", ",").split(",") if x.strip()]
        unknown = [x for x in ids if x not in valid]
        if unknown:
            bad.append((pid, unknown))
            continue
        got[pid] = sorted(set(ids))
    if bad:
        raise SystemExit("unknown skill ids (fix the sheet, nothing was imported): " + "; ".join(f"{p}: {u}" for p, u in bad[:10]))
    with SessionLocal() as s:
        for pid, ids in got.items():
            old = s.scalars(select(Annotation).where(Annotation.posting_id == pid, Annotation.annotator == name)).first()
            if old:
                old.skill_ids = ids
            else:
                s.add(Annotation(posting_id=pid, annotator=name, skill_ids=ids))
        s.commit()
    print(f"{name}: {len(got)} postings imported, {empty} left empty")


def main() -> None:
    init_db()
    cmd = sys.argv[1] if len(sys.argv) > 1 else "status"
    args = sys.argv[2:]
    if cmd == "sheet" and args:
        return cmd_sheet(args[0], args[1] if len(args) > 1 else None)
    if cmd == "import-sheet" and len(args) == 2:
        return cmd_import_sheet(args[0], args[1])
    {"split": cmd_split, "split-bisnis": cmd_split_bisnis, "status": cmd_status, "consensus": cmd_consensus}.get(cmd, lambda: sys.exit(__doc__))()


if __name__ == "__main__":
    main()
