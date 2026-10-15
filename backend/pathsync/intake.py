"""Turning student uploads into evidence (paper III.A.2, Fig. 5).

transcript  -> course x grade x mapping strength       -> academic exposure
certificate -> completion level x mapping strength      -> completion evidence
project     -> rubric 0-3 per skill with verbatim quotes -> demonstrated skill
CV          -> claimed skills with quotes                -> confirmation only (not mastery)
"""
from __future__ import annotations

import hashlib
import io
import json
import re
import secrets
from functools import lru_cache

from .config import GRADE_SCORES, RUBRIC_LABELS, SEED_DIR
from .embeddings import encode
from .evidence import EvidenceItem, rubric_score
from .extraction import extract_skills, skill_index
from .extraction.llm import LLMUnavailable, llm_configured
from .taxonomy import load_skills
from .text import normalize_for_match, verify_quote

COURSE_MATCH = 0.80       # embedding similarity to accept a catalogue course
SKILL_FALLBACK = 0.60     # direct course-name -> skill similarity for unknown courses


# ------------------------------------------------------------------ documents
MAX_PDF_PAGES = 40                       # a transcript, CV or project report; more pages is abuse or the wrong file
MAX_DOCX_UNZIPPED = 40 * 1024 * 1024     # a 5 MB .docx must not expand into gigabytes (zip bomb)


class DocumentRejected(ValueError):
    """The file is readable in principle but too large or damaged to parse safely."""


def document_to_text(data: bytes, filename: str) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        import pdfplumber

        with pdfplumber.open(io.BytesIO(data)) as pdf:
            if len(pdf.pages) > MAX_PDF_PAGES:
                raise DocumentRejected(f"This PDF has {len(pdf.pages)} pages; the limit is {MAX_PDF_PAGES}.")
            return "\n".join((p.extract_text() or "") for p in pdf.pages)
    if name.endswith(".docx"):
        import zipfile

        import docx

        try:
            unzipped = sum(i.file_size for i in zipfile.ZipFile(io.BytesIO(data)).infolist())
        except zipfile.BadZipFile as e:
            raise DocumentRejected("This .docx file is damaged or is not a Word document.") from e
        if unzipped > MAX_DOCX_UNZIPPED:
            raise DocumentRejected("This .docx expands to more than 40 MB. Save it again, or paste the text instead.")
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    if name.endswith(".ipynb"):
        nb = json.loads(data.decode("utf-8", errors="replace"))
        cells = []
        for c in nb.get("cells", []):
            src = "".join(c.get("source", []))
            cells.append(src)
        return "\n\n".join(cells)
    return data.decode("utf-8", errors="replace")


# ------------------------------------------------------------------ catalogue
@lru_cache(maxsize=1)
def _courses() -> list[dict]:
    return json.loads((SEED_DIR / "courses.json").read_text(encoding="utf-8"))["courses"]


@lru_cache(maxsize=1)
def _certs() -> dict:
    return json.loads((SEED_DIR / "certificates.json").read_text(encoding="utf-8"))


def _catalog_index(entries: list[dict]):
    forms, owner = [], []
    for i, e in enumerate(entries):
        for f in [e["name"], *e.get("aliases", [])]:
            forms.append(f)
            owner.append(i)
    return forms, owner, encode(forms)


@lru_cache(maxsize=1)
def _course_index():
    return _catalog_index(_courses())


@lru_cache(maxsize=1)
def _cert_index():
    return _catalog_index(_certs()["certificates"])


def _match(name: str, entries: list[dict], index) -> tuple[dict | None, float, str]:
    norm = normalize_for_match(name)
    for e in entries:
        if norm in {normalize_for_match(f) for f in [e["name"], *e.get("aliases", [])]}:
            return e, 1.0, e["name"]
    forms, owner, mat = index
    sims = (encode([name]) @ mat.T)[0]
    j = int(sims.argmax())
    return (entries[owner[j]], float(sims[j]), forms[j]) if sims[j] >= COURSE_MATCH else (None, float(sims[j]), forms[j])


def map_course(name: str) -> dict:
    entry, sim, form = _match(name, _courses(), _course_index())
    if entry is not None:
        factor = 1.0 if sim >= 0.999 else round(sim, 3)
        return {"input": name, "matched": entry["name"], "similarity": round(sim, 3), "method": "katalog",
                "skills": {k: round(v * factor, 3) for k, v in entry["skills"].items()}}
    tops = [(s, v) for s, v in skill_index().top_k(name, 3) if v >= SKILL_FALLBACK][:2]
    return {"input": name, "matched": None, "similarity": round(sim, 3), "method": "embedding" if tops else "tidak_terpetakan",
            "skills": {s: round(min(0.6, v - 0.1), 3) for s, v in tops}}


GRADE_RE = re.compile(r"(?<![A-Za-z])(A-|A|B\+|B-|B|C\+|C|D|E)(?![A-Za-z+\-])")


def parse_transcript(text: str) -> list[dict]:
    """Parse 'nama mata kuliah ... nilai' lines from pasted text, CSV or a PDF transcript."""
    out = []
    for raw in text.splitlines():
        line = raw.strip()
        if len(line) < 4:
            continue
        if "," in line or ";" in line or "\t" in line:
            parts = [p.strip() for p in re.split(r"[,;\t]", line) if p.strip()]
            grades = [p for p in parts if p.upper() in GRADE_SCORES]
            names = [p for p in parts if p.upper() not in GRADE_SCORES and not re.fullmatch(r"[\d.]+|[A-Z]{2,4}\d{3,6}", p)]
            if grades and names:
                out.append({"name": max(names, key=len), "grade": grades[-1].upper()})
            continue
        m = list(GRADE_RE.finditer(line))
        if not m:
            continue
        g = m[-1]
        name = line[:g.start()]
        name = re.sub(r"\b[A-Z]{2,5}\d{3,6}\b", "", name)          # course codes
        name = re.sub(r"\b\d+(\.\d+)?\s*(sks|SKS|credits?)?\b", "", name)
        name = re.sub(r"\s{2,}", " ", name).strip(" -:|.")
        if len(name) >= 4 and not name.lower().startswith(("ipk", "gpa", "total", "semester")):
            out.append({"name": name, "grade": g.group(1).upper()})
    return out


def academic_evidence(courses: list[dict]) -> tuple[EvidenceItem, list[dict]]:
    scores: dict[str, float] = {}
    strength: dict[str, float] = {}
    detail = []
    for c in courses:
        grade = str(c.get("grade", "")).upper().strip()
        if grade not in GRADE_SCORES:
            continue
        m = map_course(c["name"])
        gs = GRADE_SCORES[grade]
        contrib = {}
        for s, st in m["skills"].items():
            e = gs * st
            contrib[s] = round(e, 3)
            scores[s] = max(scores.get(s, 0.0), e)
            strength[s] = max(strength.get(s, 0.0), st)
        detail.append({**m, "grade": grade, "grade_score": gs, "contribution": contrib})
    return EvidenceItem("academic", "Transcript", {k: round(v, 4) for k, v in scores.items()}, strength), detail


def _english_score(scale: str | None, score: float | None) -> float:
    if not scale or score is None:
        return 0.75
    if scale == "toeic":
        return 0.9 if score >= 785 else 0.75 if score >= 605 else 0.55
    if scale == "toefl":
        return 0.9 if score >= 550 else 0.75 if score >= 500 else 0.55
    if scale == "ielts":
        return 0.9 if score >= 6.5 else 0.75 if score >= 5.5 else 0.55
    return 0.75


VERIFY_DOMAINS = ("coursera.org", "credly.com", "dicoding.com", "linkedin.com", "grow.google", "skillshop.exceedlms.com",
                  "learn.microsoft.com", "aws.amazon.com", "kaggle.com", "datacamp.com", "udemy.com", "hackerrank.com",
                  "edx.org", "freecodecamp.org", "accredible.com", "bnsp.go.id", "cloudskillsboost.google", "ets.org",
                  "myskills.id", "revou.co", "skilvul.com", "mysertifikasi.com")


def verification_link(url: str | None) -> str | None:
    """A credential URL on a known issuer's domain, or None. The page itself is not fetched."""
    from urllib.parse import urlparse

    u = (url or "").strip()
    if not u:
        return None
    if "://" not in u:
        u = "https://" + u
    try:
        host = (urlparse(u).hostname or "").lower()
    except ValueError:
        return None
    return u[:500] if any(host == d or host.endswith("." + d) for d in VERIFY_DOMAINS) else None


def certificate_evidence(certs: list[dict]) -> tuple[EvidenceItem, list[dict]]:
    levels = _certs()["levels"]
    scores, strength, detail = {}, {}, []
    for c in certs:
        entry, sim, form = _match(c["name"], _certs()["certificates"], _cert_index())
        if entry is not None:
            lvl = levels[entry["level"]]
            mapping = entry["skills"]
            method = "katalog"
        else:
            lvl = levels["kursus"]
            mapping = {s: round(min(0.6, v - 0.1), 3) for s, v in skill_index().top_k(c["name"], 2) if v >= SKILL_FALLBACK}
            method = "embedding" if mapping else "tidak_terpetakan"
        contrib = {}
        for s, st in mapping.items():
            e = lvl * st
            if entry is not None and entry.get("score_scale") and s == "english_proficiency":
                e = _english_score(entry["score_scale"], c.get("score"))
            contrib[s] = round(e, 3)
            scores[s] = max(scores.get(s, 0.0), e)
            strength[s] = max(strength.get(s, 0.0), st)
        detail.append({"input": c["name"], "matched": entry["name"] if entry else None, "similarity": round(sim, 3),
                       "method": method, "level": entry["level"] if entry else "kursus", "contribution": contrib,
                       "url": verification_link(c.get("url")), "verified_link": bool(verification_link(c.get("url")))})
    return EvidenceItem("certificate", "Certificates", {k: round(v, 4) for k, v in scores.items()}, strength), detail


def cv_evidence(text: str) -> tuple[EvidenceItem, dict]:
    res = extract_skills(text, kind="cv")
    claims = sorted(res.skills)
    detail = {"extractor": res.extractor,
              "claims": [{"skill_id": s, "quote": m.quote, "level": m.level, "surface": m.surface} for s, m in res.skills.items()],
              "new_candidates": [m.surface for m in res.new_candidates][:20]}
    return EvidenceItem("cv", "CV", {}, {}, claims=claims), detail


# ---------------------------------------------------------------- project rubric
# Without an LLM rubric or a human reviewer the heuristic rubric never exceeds "mandiri" (PM audit P0-2):
# keywords and code snippets can be copied from any tutorial.
HEURISTIC_MAX_LEVEL = 2
_SQL_L1 = re.compile(r"\bselect\b[\s\S]{1,200}?\bfrom\b", re.I)
_SQL_L2 = re.compile(r"\b(join|group\s+by|having)\b", re.I)
_SQL_L3 = re.compile(r"\bover\s*\(|\bwith\s+\w+\s+as\s*\(|\bpartition\s+by\b|\bcase\s+when\b", re.I)
_PY_L1 = re.compile(r"import\s+pandas|\bpd\.read_|import\s+numpy", re.I)
_PY_L2 = re.compile(r"\.groupby\(|\.merge\(|\.pivot_table\(|\.agg\(", re.I)
_PY_L3 = re.compile(r"\bdef\s+\w+\(|from\s+sklearn|statsmodels|scipy\.stats", re.I)
_VIZ = re.compile(r"plt\.|sns\.|plotly|\.plot\(|chart|grafik|visualisasi|dashboard", re.I)


def _heuristic_rubric(text: str) -> list[dict]:
    """Transparent fallback when no LLM is configured: level from depth of evidence, always quoted."""
    res = extract_skills(text, prefer="lexicon")
    by_skill: dict[str, list] = {}
    for m in res.mentions:
        by_skill.setdefault(m.skill_id, []).append(m)
    ratings = []
    for s, ms in by_skill.items():
        n_quotes = len({m.quote for m in ms})
        # keyword mentions alone show "dasar"; repeated mentions, or a mention inside a sentence describing work the
        # student did, show "mandiri". "mahir" is never given without an LLM rubric, a human reviewer or an assessment.
        worked = [m for m in ms if _WORK_CUES.search(m.quote)]
        level = 2 if (n_quotes >= 2 or worked) else 1
        quote = _focus(text, (worked or ms)[0].quote, (worked or ms)[0].surface)
        rationale = f"{n_quotes} part{'s' if n_quotes > 1 else ''} of the artifact show{'' if n_quotes > 1 else 's'} this skill."
        if s == "sql":
            level = 3 if _SQL_L3.search(text) else 2 if _SQL_L2.search(text) else 1 if _SQL_L1.search(text) else min(level, 1)
            m = _SQL_L3.search(text) or _SQL_L2.search(text) or _SQL_L1.search(text)
            if m:
                quote = _line_at(text, m.start())
            rationale = "Query depth: " + {1: "basic SELECT", 2: "JOIN/aggregation", 3: "window functions/CTEs"}.get(level, "-")
        elif s == "python_analytics":
            level = 3 if _PY_L3.search(text) else 2 if _PY_L2.search(text) else 1 if _PY_L1.search(text) else min(level, 1)
            m = _PY_L3.search(text) or _PY_L2.search(text) or _PY_L1.search(text)
            if m:
                quote = _line_at(text, m.start())
        ratings.append({"skill_id": s, "level": min(level, HEURISTIC_MAX_LEVEL), "quote": quote, "rationale": rationale})
    rated = {r["skill_id"] for r in ratings}
    if "sql" not in rated and _SQL_L1.search(text):
        lvl = 3 if _SQL_L3.search(text) else 2 if _SQL_L2.search(text) else 1
        m = _SQL_L3.search(text) or _SQL_L2.search(text) or _SQL_L1.search(text)
        ratings.append({"skill_id": "sql", "level": min(lvl, HEURISTIC_MAX_LEVEL), "quote": _line_at(text, m.start()),
                        "rationale": "Query depth: " + {1: "basic SELECT", 2: "JOIN/aggregation", 3: "window functions/CTEs"}[lvl]})
    if "python_analytics" not in rated and _PY_L1.search(text):
        lvl = 3 if _PY_L3.search(text) else 2 if _PY_L2.search(text) else 1
        m = _PY_L3.search(text) or _PY_L2.search(text) or _PY_L1.search(text)
        ratings.append({"skill_id": "python_analytics", "level": min(lvl, HEURISTIC_MAX_LEVEL), "quote": _line_at(text, m.start()),
                        "rationale": "Data-analysis code in Python."})
    if _VIZ.search(text) and "data_visualization" not in by_skill:
        m = _VIZ.search(text)
        ratings.append({"skill_id": "data_visualization", "level": 1, "quote": _line_at(text, m.start()),
                        "rationale": "The artifact builds charts."})
    return ratings


def _line_at(text: str, pos: int, width: int = 220) -> str:
    """The line containing `pos`; for long lines (PDF text often has none) a window around `pos`."""
    a = text.rfind("\n", 0, pos) + 1
    b = text.find("\n", pos)
    b = b if b != -1 else len(text)
    if b - a <= width:
        return text[a:b].strip()
    a = max(a, pos - width // 3)
    return text[a: min(b, a + width)].strip()


def _focus(text: str, quote: str, surface: str) -> str:
    """Shorten a long quote to the part around the matched term, so the evidence points at the relevant line."""
    if len(quote) <= 220:
        return quote
    i = text.find(quote)
    j = quote.lower().find((surface or "").lower())
    if i >= 0 and j >= 0:
        return _line_at(text, i + j)
    return quote[:220]


# ------------------------------------------------------------- artifact integrity
_POSTING_CUES = re.compile(
    r"\b(we are looking for|you will|requirements?|qualifications?|kualifikasi|persyaratan|minimal \d|minimum \d|"
    r"years? of experience|tahun pengalaman|apply now|lamar|benefits?|tanggung jawab|job description|deskripsi pekerjaan|"
    r"responsibilities|nice to have|preferred|diutamakan|kandidat|candidates?)\b", re.I)
_WORK_CUES = re.compile(
    r"\b(saya|kami|aku|i|we|my|our|built|created|designed|analy[sz]ed|wrote|implemented|membuat|merancang|menganalisis|"
    r"mengolah|menulis|mengembangkan|hasil|result|findings?|temuan|insight|kesimpulan|conclusion)\b|select\s|import\s|def\s",
    re.I)


# lines that state a requirement ("Mampu ...", "Minimal 2 tahun ...", "Experience in ...") - typical of postings,
# rare in a student's own report, notebook or query file
_REQ_LINE = re.compile(
    r"^\s*(?:[-*•\d.)]+\s*)?(mampu|memiliki|minimal|min\.|menguasai|berpengalaman|pengalaman|pendidikan|"
    r"lulusan|bersedia|diutamakan|familiar|able to|ability to|experience|experienced|strong|proficient|proficiency|"
    r"bachelor|degree|knowledge of|good|excellent|understanding of|must|willing|at least|fresh graduate|"
    r"you will|you'll|responsible for|bertanggung jawab|melakukan|menyusun|membantu|support)\b", re.I)
_POSTING_HEADING = re.compile(
    r"(^|\n)\s*(kualifikasi|persyaratan|requirements?|qualifications?|responsibilities|tanggung jawab|"
    r"job description|deskripsi pekerjaan|what you.ll do|who you are|benefits?)\s*:?\s*(\n|$)", re.I)


@lru_cache(maxsize=1)
def _posting_corpus():
    """TF-IDF index over scraped job postings, to reject a posting pasted as 'my project'."""
    from sklearn.feature_extraction.text import TfidfVectorizer

    from .config import RAW_DIR

    texts = []
    for f in sorted(RAW_DIR.glob("*.jsonl")):
        for line in f.read_text(encoding="utf-8").splitlines():
            if line.strip():
                texts.append(json.loads(line).get("description_text", "")[:4000])
    if not texts:
        return None, None
    vec = TfidfVectorizer(analyzer="char_wb", ngram_range=(3, 5), min_df=1, sublinear_tf=True)
    return vec, vec.fit_transform(texts)


def artifact_checks(text: str) -> dict:
    """Signals that an upload is not the student's own work (paper III.A.2 requires quoted, checkable evidence)."""
    from sklearn.metrics.pairwise import cosine_similarity

    flags, sims = [], 0.0
    vec, mat = _posting_corpus()
    if vec is not None:
        sims = float(cosine_similarity(vec.transform([text[:4000]]), mat).max())
        if sims >= 0.55:
            flags.append("copy_of_job_posting")
    lines = [ln for ln in re.split(r"[\n.;]+", text) if ln.strip()]
    words = re.findall(r"\w+", text)
    posting_share = sum(bool(_POSTING_CUES.search(ln)) for ln in lines) / max(1, len(lines))
    work_share = sum(bool(_WORK_CUES.search(ln)) for ln in lines) / max(1, len(lines))
    alias_hits = len(lexicon_hits(text))
    density = alias_hits / max(1, len(words))
    req_lines = sum(bool(_REQ_LINE.match(ln)) for ln in lines) / max(1, len(lines))
    heading = bool(_POSTING_HEADING.search(text))
    if (posting_share >= 0.25 and work_share < 0.15) or req_lines >= 0.3 or (heading and req_lines >= 0.12):
        flags.append("reads_like_job_requirements")
    if density >= 0.18 or (len(words) / max(1, len(lines)) < 4 and alias_hits >= 8):
        flags.append("keyword_list")
    if len(words) < 60:
        flags.append("too_short")
    return {"flags": flags, "posting_similarity": round(sims, 3), "posting_cue_share": round(posting_share, 3),
            "work_statement_share": round(work_share, 3), "skill_term_density": round(density, 3)}


# ------------------------------------------------------------- duplicate uploads (PM audit P0-2)
_SIG_SIZE = 64
_SIG_PRIME = (1 << 61) - 1


def artifact_signature(text: str) -> list[int]:
    """MinHash over word 5-shingles: two uploads with Jaccard similarity J share about J of their signature values."""
    words = re.findall(r"\w+", text.lower())
    shingles = {" ".join(words[i:i + 5]) for i in range(max(1, len(words) - 4))} or {""}
    hashes = [int.from_bytes(hashlib.blake2b(sh.encode("utf-8"), digest_size=8).digest(), "big") for sh in shingles]
    sig = []
    for k in range(_SIG_SIZE):
        a, b = 2 * k + 1_000_003, 7919 * k + 17
        sig.append(min((a * h + b) % _SIG_PRIME for h in hashes))
    return sig


def signature_similarity(a: list[int], b: list[int]) -> float:
    if not a or not b or len(a) != len(b):
        return 0.0
    return sum(x == y for x, y in zip(a, b, strict=True)) / len(a)


def lexicon_hits(text: str) -> list[str]:
    from .extraction import lexicon_extractor

    return [m.group(1) for m in lexicon_extractor().alias_re.finditer(text)]


def project_evidence(text: str, title: str, candidate_ids: list[str] | None = None) -> tuple[EvidenceItem, dict]:
    checks = artifact_checks(text)
    if "copy_of_job_posting" in checks["flags"] or "reads_like_job_requirements" in checks["flags"]:
        detail = {"method": "rejected", "ratings": [], "rejected": [], "checks": checks, "review_status": "rejected",
                  "message": "This looks like a job posting or a list of requirements, not work you did. "
                             "Upload your own report, notebook, query file or document."}
        return EvidenceItem("project", title or "Project", {}, {}), detail
    method = "heuristik"
    ratings: list[dict] = []
    if llm_configured():
        try:
            from .extraction import llm_extractor

            cands = candidate_ids or sorted({m.skill_id for m in extract_skills(text, prefer="lexicon").mentions})
            if cands:
                ratings = llm_extractor().rate_project(text, cands)
                method = "llm"
        except LLMUnavailable:
            ratings = []
    if not ratings:
        ratings = _heuristic_rubric(text)
        method = "heuristik"
    skills = load_skills()
    scores, strength, accepted, rejected = {}, {}, [], []
    for r in ratings:
        if r["skill_id"] not in skills or int(r["level"]) <= 0:
            continue
        if not verify_quote(r.get("quote", ""), text):        # quote must come from the artifact
            rejected.append({**r, "reason": "quote not found in the artifact"})
            continue
        e = rubric_score(int(r["level"]))
        scores[r["skill_id"]] = max(scores.get(r["skill_id"], 0.0), e)
        strength[r["skill_id"]] = 1.0
        accepted.append({**r, "label": RUBRIC_LABELS[int(r["level"])], "score": e})
    # paper: human re-check on a 20% sample of rubric ratings; suspicious uploads always go to review
    sampled = secrets.randbelow(5) == 0      # truly random, so re-titling an upload cannot dodge the sample
    status = "pending" if (sampled or checks["flags"]) else "auto"
    detail = {"method": method, "ratings": accepted, "rejected": rejected, "checks": checks,
              "needs_human_review": status == "pending", "review_status": status,
              "review_reason": ", ".join(checks["flags"]) or ("random 20% quality sample" if sampled else "")}
    item = EvidenceItem("project", title or "Project", scores, strength, reliability=0.5 if status == "pending" else 1.0)
    return item, detail
